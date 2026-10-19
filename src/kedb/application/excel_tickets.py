"""Excel template and validation for resolved-ticket batch uploads."""
from io import BytesIO
import json
from uuid import NAMESPACE_URL, uuid5

COLUMNS = ("subject", "description", "resolution")


def generated_ticket_key(summary, description, resolution):
    content = json.dumps([summary.strip(), description.strip(), resolution.strip()], ensure_ascii=False)
    return "KEDB-" + str(uuid5(NAMESPACE_URL, "kedb:manual:" + content))


def ticket_template():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Tickets"
    sheet.append(COLUMNS)
    sheet.freeze_panes = "A2"
    for column, width in zip("ABC", (45, 65, 65)):
        sheet.column_dimensions[column].width = width
    instructions = book.create_sheet("Instructions")
    for line in (
        "Fill in the Tickets sheet, one resolved ticket per row.",
        "All three columns are required. Keep the header names unchanged.",
        "Ticket IDs are generated automatically; do not add an ID column.",
        "subject: a short summary of the incident.",
        "description: symptoms and relevant error messages.",
        "resolution: the actual fix that resolved the incident.",
        "Tickets are treated as Resolved and processed using the normal review rules.",
        "Identical ticket content reuses the same ID to avoid duplicate processing.",
    ):
        instructions.append([line])
    instructions.column_dimensions["A"].width = 100
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def read_ticket_workbook(content):
    from openpyxl import load_workbook
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("Cannot read this workbook. Upload a valid .xlsx file using the template.") from exc
    try:
        if "Tickets" not in book.sheetnames:
            raise ValueError("The workbook must contain a Tickets sheet.")
        rows = book["Tickets"].iter_rows()
        header = next(rows, ())
        names = [str(cell.value or "").strip() for cell in header]
        if any(names.count(name) != 1 for name in COLUMNS):
            raise ValueError("Include each required column once: " + ", ".join(COLUMNS))
        positions = [names.index(name) for name in COLUMNS]
        tickets = []
        for number, cells in enumerate(rows, 2):
            if all(cell.value is None or str(cell.value).strip() == "" for cell in cells):
                continue
            selected = [cells[index] for index in positions]
            if any(cell.data_type in {"f", "e"} for cell in selected):
                raise ValueError(f"Row {number}: replace formulas or Excel errors with text values.")
            values = [str(cell.value).strip() if cell.value is not None else "" for cell in selected]
            if not all(values):
                raise ValueError(f"Row {number}: all three ticket fields are required.")
            tickets.append(dict(zip(COLUMNS, values), status="Resolved",
                                ticket_id=generated_ticket_key(*values)))
        if not tickets:
            raise ValueError("Add at least one ticket to the Tickets sheet.")
        return tickets
    finally:
        book.close()
